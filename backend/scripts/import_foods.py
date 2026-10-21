"""Import (or update) the food database: USDA FoodData Central and the Indian
Nutrient Databank (INDB).

    cd backend && uv run --group import python -m scripts.import_foods

- USDA FoodData Central (https://fdc.nal.usda.gov), SR Legacy (2018, final) and
  Foundation Foods: public domain (CC0). Ingredients, nutrients per 100 g.
- INDB (https://github.com/lindsayjaacks/Indian-Nutrient-Databank-INDB-; Vijayakumar
  et al., Current Developments in Nutrition, 2024): about 1,000 Indian recipes built
  on the Indian Food Composition Tables (ICMR-NIN). Dishes, nutrients per serving,
  because INDB's per-100 g values include cooking water.

Dishes get a meal role (breakfast, main, staple, …) from their book chapter and
name, a diet (vegetarian / eggetarian / non-vegetarian) from their ingredients, and
pass data checks (protein, carbs and fat must add up to the calories; serving sizes
must be plausible); a dish that fails is kept for logging but never planned.

Uses SUPABASE_SECRET_KEY (backend/.env), because foods are read-only for everyone else.
"""

import csv
import io
import json
import re
import sys
import urllib.request
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from tempfile import gettempdir
from typing import Any

from dotenv import dotenv_values
from openpyxl import load_workbook

USDA_BASE = "https://fdc.nal.usda.gov/fdc-datasets/"
USDA_FILES = [
    "FoodData_Central_sr_legacy_food_csv_2018-04.zip",
    "FoodData_Central_foundation_food_csv_2026-04-30.zip",
]
# Pin a version: change this (and re-run) to pick up the databank's updates.
INDB_COMMIT = "fbdb62bec519c6e2468582a799fb47e194b42558"
INDB_BASE = (
    f"https://raw.githubusercontent.com/lindsayjaacks/Indian-Nutrient-Databank-INDB-/{INDB_COMMIT}/"
)
CACHE = Path(gettempdir()) / "workout-strategy-foods"
BATCH = 500

# --- Nutrients -------------------------------------------------------------------------

# USDA nutrient id → our key (per 100 g).
USDA_NUTRIENTS = {
    1003: "protein_g",
    1004: "fat_g",
    1005: "carbs_g",
    1079: "fibre_g",
    1087: "calcium_mg",
    1089: "iron_mg",
    1090: "magnesium_mg",
    1092: "potassium_mg",
    1093: "sodium_mg",
    1095: "zinc_mg",
    1106: "vitamin_a_ug",
    1114: "vitamin_d_ug",
    1162: "vitamin_c_mg",
    1177: "folate_ug",
    1178: "b12_ug",
    2000: "sugar_g",
}
# Energy: SR Legacy uses 1008; Foundation Foods often only the Atwater factors.
USDA_ENERGY = (1008, 2048, 2047)
MACROS = ("protein_g", "carbs_g", "fat_g", "fibre_g")

# INDB column (per serving, after "unit_serving_") → our key.
INDB_MICROS = {
    "calcium_mg": "calcium_mg",
    "iron_mg": "iron_mg",
    "magnesium_mg": "magnesium_mg",
    "potassium_mg": "potassium_mg",
    "sodium_mg": "sodium_mg",
    "zinc_mg": "zinc_mg",
    "vita_ug": "vitamin_a_ug",
    "vitc_mg": "vitamin_c_mg",
    "folate_ug": "folate_ug",
    "freesugar_g": "sugar_g",
}

# --- Diet --------------------------------------------------------------------------------

MEAT = re.compile(
    r"\b(chicken|mutton|lamb|beef|pork|veal|goat|turkey|duck|fish|prawns?|shrimps?|crab|"
    r"lobster|squid|mackerel|pomfret|rohu|salmon|tuna|sardines?|anchov\w*|bacon|ham|salami|"
    r"sausages?|keema|meat|liver|gelatine?)\b",
    re.IGNORECASE,
)
EGG = re.compile(r"\b(eggs?|omelette|omlet|mayonnaise)\b", re.IGNORECASE)
# USDA categories that are always meat or fish.
USDA_MEAT_CATEGORIES = {5, 7, 10, 13, 15, 17}
# Left out: baby foods, regional native foods, branded and lab samples.
USDA_SKIP_CATEGORIES = {3, 24, 26, 27}

# --- Common foods the planner uses (USDA) ----------------------------------------------
# fdc id → (plain name, role, serving label, grams per serving, countable)
COMMON_USDA = {
    "173424": ("Boiled eggs", "protein", "egg", 50, True),
    "171477": ("Grilled chicken breast", "protein", "100 g", 100, False),
    "173709": ("Tuna (canned in water)", "protein", "100 g", 100, False),
    "171284": ("Curd (plain yogurt)", "protein", "bowl", 150, False),
    "170894": ("Greek yogurt (plain)", "protein", "cup", 170, False),
    "172448": ("Tofu (firm)", "protein", "100 g", 100, False),
    "173757": ("Boiled chickpeas", "protein", "bowl", 150, False),
    "170870": ("Milk (low-fat)", "drink", "glass", 250, False),
    "169957": ("Moong sprouts", "snack", "bowl", 100, False),
    "173806": ("Roasted peanuts", "snack", "handful", 30, False),
    "170567": ("Almonds", "snack", "handful", 25, False),
    "170187": ("Walnuts", "snack", "handful", 20, False),
    "173944": ("Banana", "fruit", "banana", 118, True),
    "171688": ("Apple", "fruit", "apple", 180, True),
    "169097": ("Orange", "fruit", "orange", 130, True),
    "169926": ("Papaya", "fruit", "bowl", 150, False),
    "173044": ("Guava", "fruit", "guava", 100, True),
    "169134": ("Pomegranate", "fruit", "bowl", 100, False),
    "174683": ("Grapes", "fruit", "bowl", 100, False),
    "167765": ("Watermelon", "fruit", "bowl", 200, False),
    "169910": ("Mango", "fruit", "bowl", 150, False),
    "167776": ("Pear", "fruit", "pear", 180, True),
    "169704": ("Brown rice (cooked)", "staple", "bowl", 150, False),
    "172688": ("Whole-wheat bread", "staple", "slice", 32, True),
}

# --- INDB meal roles -----------------------------------------------------------------------

# The Art & Science of Cooking: chapter → role ('skip': stocks, preserves and the like).
ASC_CHAPTERS = {
    "5": "drink", "6": "breakfast", "7": "breakfast", "8": "skip", "9": "staple",
    "10": "main", "11": "main", "12": "main", "13": "salad", "14": "sweet", "15": "sweet",
    "16": "snack", "17": "sweet", "18": "sweet", "19": "sweet", "20": "snack", "21": "snack",
    "22": "skip", "23": "skip", "24": "skip", "25": "skip",
}  # fmt: skip
ROLE_WORDS = [
    ("skip", r"stock|garam masala|chat masala|chaat masala|kashmiri masala|powder|squash|jam|"
             r"jelly|murabba|pickle|achaar|chutney|sauce|puree|premix|consomm\w*|dough|pastry|"
             r"syrup|ketchup|mayonnaise|dressing|paste|icing|"
             r"filling|baghar|tadka|preserves|pickled|spice blend|infant|shishu|amylase|weaning|"
             r"cereal mix|pulse mix|vegetable mix"),
    ("drink", r"tea|coffee|chai|lassi|milkshake|shake|juice|sharbat|sherbet|punch|cooler|"
              r"smoothie|panna|jal jeera|mintade|lemonade|kehwa|buttermilk|chaas|thandai|kanji|"
              r"canjee|hot chocolate|doodh|sattu drink"),
    ("soup", r"soup|rasam|shorba"),
    ("sweet", r"kheer|halwa|burfi|barfi|ladoos?|laddoos?|laddus?|cake|pudding|custard|"
              r"payasam|gulab|jalebi|rasgulla|sandesh|ice cream|kulfi|mousse|tart|pie|souffle|"
              r"trifle|phirni|"
              r"shrikhand|peda|mithai|sweet|gateau|flan|pavlova|fool|whip|chum chum|rasmalai|"
              r"rabri|malpua|chikki|cheesecake|sunday|bavarian|cream horns|cream buns|shells|"
              r"swans|kesari|swiss roll|cookies?|biscuits?|ginger bread|brownies?|muffins?"),
    ("breakfast", r"poha|upma|idli|dosa|uttapam|uthappam|cheela|chilla|omelette|omlet|porridge|"
                  r"daliya|dalia|oats|oatmeal|cornflakes|flakes|sandwich|toast|pancake|bhujia|"
                  r"paniyaram|appam|pongal|thepla|muesli|scrambled|boiled egg|poached|handvo"),
    ("staple", r"roti|chapati|chapatti|parantha|paratha|poori|puri|naan|kulcha|bhatura|bhature|"
               r"rice|pulao|pulav|biryani|khichdi|khichri|bread|phulka|missi|makki|bhakri|"
               r"rotla|tahar"),
    ("salad", r"salad|raita|kachumber|koshimbir"),
    ("snack", r"pakora|pakoda|bhajji|bajji|samosa|kachori|tikki|cutlet|vada|dhokla|chaat|chat|"
              r"bhel|sev|namak paras?|mathri|muthias?|kebab|kabab|tikka|roll|sundal|makhana|"
              r"cookies|biscuit|nuts|murmura|chiwda|chivda|mixture|khakhra|straws|puffs|"
              r"aigrettes|patties|horns|chips|namkeen|papdi|gobi 65"),
    ("main", r"dal|daal|curry|sabzi|subzi|bhaji|paneer|chicken|mutton|fish|prawn|egg|keema|"
             r"kofta|korma|sambar|kadhi|rajma|chole|chana|aloo|gobhi|matar|bhindi|baingan|"
             r"palak|saag|stew|dalma|avial|thoran|poriyal|kootu|usal|jhol|vindaloo|do pyaza|"
             r"bharta|bhartha|methi|lauki|ghiya|tinda|karela|karele|mushroom|soyabean|soya|tofu|"
             r"josh|yakhni|gosht|meat|kurma|stir|fry|fried|urad|arbi|foogath|jackfruit|kathal|"
             r"jalfrezi|broccoli|dhansak|cauliflower|potato|potatoes|au gratin|spinach|"
             r"manchurian|beans|macaroni|spaghetti|lasagne|noodles|pizza|sajina|radishes"),
]  # fmt: skip
ROLE_PATTERNS = [(role, re.compile(rf"\b({words})\b")) for role, words in ROLE_WORDS]
# Deep-fried or very rich: INDB counts all the frying oil, and they don't suit a
# health plan anyway. Kept for logging, never planned.
FRIED = re.compile(
    r"\b(pakora|pakoda|pakodies|bhajji|bajji|samosa|kachori|poori|puri|bhatura|bhature|"
    r"fried|fritters?|chips|namak paras?|mathri|vada|malpua|jalebi|kofta|aigrettes|murukku)\b"
)
# Plausible calories for one serving, by role.
ROLE_KCAL = {
    "breakfast": (60, 700), "main": (60, 700), "staple": (50, 600), "salad": (20, 400),
    "snack": (30, 450), "soup": (20, 400), "drink": (0, 400), "sweet": (50, 500),
}  # fmt: skip
# More than this in one "piece" means the piece label can't be trusted.
PIECE_KCAL_MAX = 250
UNCOUNTABLE = {
    "bowl", "plate", "cup", "glass", "tall glass", "small bowl", "soup bowl", "dish",
    "tablespoon", "jar", "curry bowl", "katori", "souffle dish", "tea cup", "mug", "serving",
    "portion", "teaspoon", "medium bowl", "big bowl", "large bowl",
}  # fmt: skip
# Book code → grocery group (Indian Food Composition Tables food groups).
GROCERY_GROUPS = {
    "A": "Grains and flours", "B": "Pulses", "C": "Leafy greens", "D": "Vegetables",
    "E": "Fruits", "F": "Roots and tubers", "G": "Spices and herbs", "H": "Nuts and seeds",
    "I": "Sugar and sweeteners", "J": "Vegetables", "K": "Other", "L": "Dairy",
    "M": "Eggs", "N": "Meat and fish", "O": "Meat and fish", "P": "Meat and fish",
    "Q": "Meat and fish", "R": "Meat and fish", "S": "Meat and fish", "T": "Oils and fats",
    "U": "Bakery", "V": "Drinks", "W": "Other", "X": "Sauces",
}  # fmt: skip
# Rough grams for kitchen measures, so amounts add up on the grocery list.
UNIT_GRAMS = {"tsp": 5, "tbsp": 15, "c": 240}
SKIP_INGREDIENTS = re.compile(r"^(water|salt)\b", re.IGNORECASE)
# Fresh aromatics are listed with spices in the composition tables; shoppers
# look for them with the vegetables.
FRESH = re.compile(
    r"^(onions?|garlic|ginger|green chill\w*|coriander leaves|mint leaves|curry leaves|"
    r"spring onions?|lemons?)\b",
    re.IGNORECASE,
)


def download(url: str) -> bytes:
    CACHE.mkdir(exist_ok=True)
    path = CACHE / url.rsplit("/", 1)[1]
    if not path.exists():
        print(f"  downloading {path.name}")
        with urllib.request.urlopen(url, timeout=300) as response:
            path.write_bytes(response.read())
    return path.read_bytes()


def number(value: Any) -> float | None:
    try:
        n = float(value)
    except (TypeError, ValueError):
        return None
    return None if n != n else n  # NaN


def rounded(value: float | None, places: int = 2) -> float | None:
    return None if value is None else round(value, places)


# --- USDA ------------------------------------------------------------------------------------


def usda_rows() -> list[dict[str, Any]]:
    rows = []
    for name in USDA_FILES:
        archive = zipfile.ZipFile(io.BytesIO(download(USDA_BASE + name)))

        def read(member: str, archive: zipfile.ZipFile = archive) -> csv.DictReader:
            path = next(p for p in archive.namelist() if p.endswith("/" + member))
            return csv.DictReader(io.TextIOWrapper(archive.open(path), encoding="utf-8"))

        categories = {r["id"]: r["description"] for r in read("food_category.csv")}
        kind = "foundation_food" if "foundation" in name else "sr_legacy_food"
        foods = {
            r["fdc_id"]: r
            for r in read("food.csv")
            if r["data_type"] == kind
            and int(r["food_category_id"] or 0) not in USDA_SKIP_CATEGORIES
        }
        values: dict[str, dict[int, float]] = defaultdict(dict)
        wanted = set(USDA_NUTRIENTS) | set(USDA_ENERGY)
        for r in read("food_nutrient.csv"):
            nutrient = int(r["nutrient_id"])
            if r["fdc_id"] in foods and nutrient in wanted:
                amount = number(r["amount"])
                if amount is not None:
                    values[r["fdc_id"]][nutrient] = amount
        for fdc_id, food in foods.items():
            found = values.get(fdc_id, {})
            kcal = next((found[e] for e in USDA_ENERGY if e in found), None)
            if kcal is None:
                continue
            per = {USDA_NUTRIENTS[n]: v for n, v in found.items() if n in USDA_NUTRIENTS}
            category_id = int(food["food_category_id"] or 0)
            description = food["description"].strip()
            if category_id in USDA_MEAT_CATEGORIES or MEAT.search(description):
                diet = "non_vegetarian"
            elif EGG.search(description):
                diet = "eggetarian"
            else:
                diet = "vegetarian"
            common = COMMON_USDA.get(fdc_id)
            rows.append(
                {
                    "id": f"usda-{fdc_id}",
                    "name": description[:200],
                    "alias": common[0] if common else None,
                    "source": "usda",
                    "category": categories.get(food["food_category_id"]),
                    "role": common[1] if common else None,
                    "diet": diet,
                    "basis": "100g",
                    "serving_label": common[2] if common else None,
                    "serving_g": common[3] if common else None,
                    "countable": bool(common and common[4]),
                    "kcal": round(kcal, 1),
                    **{m: max(rounded(per.get(m)) or 0, 0) for m in MACROS if m != "fibre_g"},
                    "fibre_g": rounded(max(per["fibre_g"], 0)) if "fibre_g" in per else None,
                    "micros": {
                        k: rounded(v, 3)
                        for k, v in per.items()
                        if k not in MACROS and v is not None
                    },
                    "ingredients": None,
                }
            )
    return rows


# --- INDB --------------------------------------------------------------------------------------


def sheet(name: str) -> list[dict[str, Any]]:
    book = load_workbook(io.BytesIO(download(INDB_BASE + name)), read_only=True, data_only=True)
    lines = book.worksheets[0].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(lines)]
    return [dict(zip(header, line, strict=False)) for line in lines]


def indb_role(code: str, chapter: str, name: str) -> str | None:
    lower = name.lower()
    if code.startswith("ASC"):
        by_chapter = ASC_CHAPTERS.get(chapter)
        if by_chapter in ("skip", "drink"):
            return by_chapter
        # Sweets and preserves turn up in other chapters too.
        for role, pattern in ROLE_PATTERNS:
            if role in ("skip", "sweet") and pattern.search(lower):
                return role
        if by_chapter:
            return by_chapter
    return next((role for role, pattern in ROLE_PATTERNS if pattern.search(lower)), None)


def ingredient_amount(amount: Any, unit: Any) -> tuple[float, str] | None:
    value = number(amount)
    if value is None or value <= 0:
        return None
    unit = str(unit or "").strip().lower()
    if unit in ("g", "ml"):
        return value, unit
    if unit in UNIT_GRAMS:
        return value * UNIT_GRAMS[unit], "g"
    return value, unit or "nos"


def indb_rows() -> list[dict[str, Any]]:
    nutrients = sheet("INDB.xlsx")
    names = {r["recipe_code"]: r for r in sheet("recipes_names.xlsx")}
    servings = {r["recipe_code"]: r for r in sheet("recipes_servingsize.xlsx")}
    ingredients: dict[str, list[dict[str, Any]]] = defaultdict(list)
    plain_names: dict[str, Counter] = defaultdict(Counter)
    for r in sheet("recipes.xlsx"):
        if r.get("recipe_code") and r.get("food_code"):
            ingredients[r["recipe_code"]].append(r)
            plain = str(r.get("ingredient_name_org") or "").strip()
            if plain:
                plain_names[r["food_code"]][plain.capitalize()] += 1

    rows = []
    for n in nutrients:
        code = str(n["food_code"])
        name = str(n["food_name"]).strip()
        kcal = number(n.get("unit_serving_energy_kcal"))
        if not code or kcal is None:
            continue
        chapter = str((names.get(code) or {}).get("recipe_code_org") or "").split(".")[0]
        role = indb_role(code, chapter, name)
        parts = ingredients.get(code, [])
        text = " ".join([name] + [str(p.get("food_name") or "") for p in parts])
        groups = {str(p["food_code"])[:1] for p in parts}
        if groups & set("NOPQRS") or MEAT.search(text):
            diet = "non_vegetarian"
        elif "M" in groups or EGG.search(text):
            diet = "eggetarian"
        else:
            diet = "vegetarian"

        protein = number(n.get("unit_serving_protein_g")) or 0
        carbs = number(n.get("unit_serving_carb_g")) or 0
        fat = number(n.get("unit_serving_fat_g")) or 0
        unit = str(n.get("servings_unit") or "serving").strip().lower()
        countable = unit not in UNCOUNTABLE
        if countable and kcal > PIECE_KCAL_MAX:
            # INDB's "piece" is often a whole recipe portion (a 450 kcal "biscuit",
            # a 350 kcal "cheela" made from 55 g of dal): plan it as servings instead.
            countable, unit = False, "serving"
        # Data checks: macros add up to the calories (within 20%), and one serving
        # is a plausible amount for its role (small countable items may be light).
        atwater = 4 * protein + 4 * carbs + 9 * fat
        low, high = ROLE_KCAL.get(role or "", (0, 0))
        plausible = (
            kcal > 0
            and abs(atwater - kcal) / kcal < 0.2
            and (20 if countable else low) <= kcal <= high
        )
        plannable = role not in (None, "skip") and plausible and not FRIED.search(name.lower())

        per_serving = number((servings.get(code) or {}).get("no_of_servings")) or 1
        listed = []
        for p in parts:
            plain = str(p.get("ingredient_name_org") or p.get("food_name") or "")
            if SKIP_INGREDIENTS.match(plain) or SKIP_INGREDIENTS.match(str(p.get("food_name"))):
                continue
            amount = ingredient_amount(p.get("amount"), p.get("unit"))
            if not amount:
                continue
            food_code = str(p["food_code"])
            common = plain_names[food_code].most_common(1)
            listed.append(
                {
                    "code": food_code,
                    "name": common[0][0] if common else plain,
                    "amount": round(amount[0] / per_serving, 2),
                    "unit": amount[1],
                    "group": "Vegetables"
                    if FRESH.match(plain) or FRESH.match(str(p.get("food_name") or ""))
                    else GROCERY_GROUPS.get(food_code[:1], "Other"),
                }
            )
        rows.append(
            {
                "id": f"indb-{code}",
                "name": name[:200],
                "alias": None,
                "source": "indb",
                "category": role if role != "skip" else "other",
                "role": role if plannable else None,
                "diet": diet,
                "basis": "serving",
                "serving_label": unit[:40],
                "serving_g": None,
                "countable": countable,
                "kcal": round(kcal, 1),
                "protein_g": round(protein, 2),
                "carbs_g": round(carbs, 2),
                "fat_g": round(fat, 2),
                "fibre_g": rounded(number(n.get("unit_serving_fibre_g"))),
                "micros": {
                    key: round(v, 3)
                    for column, key in INDB_MICROS.items()
                    if (v := number(n.get(f"unit_serving_{column}"))) is not None
                }
                | (
                    {"vitamin_d_ug": round(d, 3)}
                    if (
                        d := (number(n.get("unit_serving_vitd2_ug")) or 0)
                        + (number(n.get("unit_serving_vitd3_ug")) or 0)
                    )
                    else {}
                ),
                "ingredients": listed,
            }
        )
    return rows


def save(url: str, key: str, rows: list[dict[str, Any]]) -> None:
    for start in range(0, len(rows), BATCH):
        batch = rows[start : start + BATCH]
        request = urllib.request.Request(
            f"{url}/rest/v1/foods?on_conflict=id",
            method="POST",
            data=json.dumps(batch).encode(),
            headers={
                "apikey": key,
                "Authorization": f"Bearer {key}",
                "Content-Type": "application/json",
                "Prefer": "resolution=merge-duplicates,return=minimal",
            },
        )
        urllib.request.urlopen(request, timeout=120).close()
        print(f"  saved {start + len(batch)}/{len(rows)}")


def main() -> None:
    env = dotenv_values(".env")
    url, key = (env.get("SUPABASE_URL") or "").rstrip("/"), env.get("SUPABASE_SECRET_KEY")
    if not url or not key:
        sys.exit("SUPABASE_URL and SUPABASE_SECRET_KEY must be set in backend/.env")
    usda = usda_rows()
    print(f"USDA: {len(usda)} foods ({sum(1 for r in usda if r['role'])} used by the planner)")
    indb = indb_rows()
    roles = Counter(r["role"] for r in indb if r["role"])
    print(f"INDB: {len(indb)} dishes; planner uses {sum(roles.values())}: {dict(roles)}")
    missing = set(COMMON_USDA) - {r["id"].removeprefix("usda-") for r in usda}
    if missing:
        sys.exit(f"Common foods not found in USDA data: {sorted(missing)}")
    save(url, key, usda + indb)
    print("Done.")


if __name__ == "__main__":
    main()
